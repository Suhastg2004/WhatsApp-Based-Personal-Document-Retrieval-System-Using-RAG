from dataclasses import dataclass
from pathlib import Path

import fitz
import pytesseract
from docx import Document as DocxDocument
from PIL import Image


@dataclass
class ParsedDocument:
    document_id: str
    title: str
    source: str
    file_type: str
    text: str


class DocumentParser:
    SUPPORTED_EXTENSIONS = {".pdf", ".txt", ".md", ".docx", ".png", ".jpg", ".jpeg", ".tif", ".tiff"}

    def __init__(self, tesseract_cmd: str | None = None) -> None:
        if tesseract_cmd:
            pytesseract.pytesseract.tesseract_cmd = tesseract_cmd

    def parse_file(self, file_path: Path, document_id: str) -> ParsedDocument:
        suffix = file_path.suffix.lower()
        if suffix not in self.SUPPORTED_EXTENSIONS:
            raise ValueError(f"Unsupported file type: {suffix}")

        if suffix in {".txt", ".md"}:
            text = file_path.read_text(encoding="utf-8", errors="ignore")
        elif suffix == ".pdf":
            text = self._parse_pdf(file_path)
        elif suffix == ".docx":
            text = self._parse_docx(file_path)
        else:
            text = self._parse_image(file_path)

        return ParsedDocument(
            document_id=document_id,
            title=file_path.stem,
            source=file_path.name,
            file_type=suffix,
            text=text,
        )

    def _parse_docx(self, file_path: Path) -> str:
        document = DocxDocument(file_path)
        return "\n".join(p.text for p in document.paragraphs if p.text.strip())

    def _parse_image(self, file_path: Path) -> str:
        image = Image.open(file_path)
        return pytesseract.image_to_string(image)

    def _parse_pdf(self, file_path: Path) -> str:
        doc = fitz.open(file_path)
        pages: list[str] = []

        for page in doc:
            page_text = page.get_text("text").strip()
            if page_text:
                pages.append(page_text)
                continue

            pix = page.get_pixmap(dpi=200)
            image = Image.frombytes("RGB", [pix.width, pix.height], pix.samples)
            ocr_text = pytesseract.image_to_string(image).strip()
            if ocr_text:
                pages.append(ocr_text)

        doc.close()
        return "\n\n".join(pages)
