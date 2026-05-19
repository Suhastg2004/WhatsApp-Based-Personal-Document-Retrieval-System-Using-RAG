"""File parsing: PDF, DOCX, plain text, and OCR for images / scanned PDFs."""
from dataclasses import dataclass
from pathlib import Path

import fitz  # PyMuPDF
import pytesseract
from docx import Document as DocxDocument
from PIL import Image

SUPPORTED_EXTENSIONS = {
    ".pdf", ".txt", ".md", ".docx",
    ".png", ".jpg", ".jpeg", ".tif", ".tiff", ".webp",
}


@dataclass
class ParsedDocument:
    title: str
    source: str
    file_type: str
    text: str


def parse_file(file_path: Path) -> ParsedDocument:
    suffix = file_path.suffix.lower()
    if suffix not in SUPPORTED_EXTENSIONS:
        raise ValueError(f"Unsupported file type: {suffix}")

    if suffix in {".txt", ".md"}:
        text = file_path.read_text(encoding="utf-8", errors="ignore")
    elif suffix == ".pdf":
        text = _parse_pdf(file_path)
    elif suffix == ".docx":
        text = _parse_docx(file_path)
    else:
        text = _parse_image(file_path)

    return ParsedDocument(
        title=file_path.stem,
        source=file_path.name,
        file_type=suffix,
        text=text,
    )


def _parse_docx(file_path: Path) -> str:
    document = DocxDocument(file_path)
    return "\n".join(p.text for p in document.paragraphs if p.text.strip())


def _parse_image(file_path: Path) -> str:
    image = Image.open(file_path)
    return pytesseract.image_to_string(image)


def _parse_pdf(file_path: Path) -> str:
    doc = fitz.open(file_path)
    pages: list[str] = []
    try:
        for page in doc:
            page_text = page.get_text("text").strip()
            if page_text:
                pages.append(page_text)
                continue
            pix = page.get_pixmap(dpi=200)
            image = Image.frombytes("RGB", [pix.width, pix.height], pix.samples)
            ocr_text = pytesseract.image_to_string(image).strip()
            if ocr_text:
                pages.append(ocr_text)
    finally:
        doc.close()
    return "\n\n".join(pages)
